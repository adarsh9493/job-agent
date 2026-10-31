#!/usr/bin/env python
"""Query and update the Excel tracker without loading the workbook into the conversation.

    python scripts/runtime/tracker.py has-applied --company "Lucid Motors" --job-id 123
    python scripts/runtime/tracker.py init
    python scripts/runtime/tracker.py summary
    python scripts/runtime/tracker.py add --company X --title Y --status Applied --match 74
    python scripts/runtime/tracker.py update --company X --job-id 123 --status Interview \
        --note "Recruiter screen booked"

Output is one line of JSON. The workbook schema is fixed (see agent/skills/tracker/SKILL.md).
"""
from __future__ import annotations

import argparse
import os
import sys
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path

from _common import TRACKER_PATH, emit, norm, today

COLUMNS = (
    "Date Applied", "Company", "Job Title", "Job ID", "Job URL", "Location",
    "Posted Date", "Fit Score", "Portal/ATS", "Status", "Notes", "Match %",
)
STATUS_COLORS = {
    "Applied": ("C6EFCE", "006100", False),
    "Interview": ("BDD7EE", "1F4E78", False),
    "Offer": ("92D050", "FFFFFF", True),
    "Rejected": ("FFC7CE", "9C0006", False),
    "Failed": ("C00000", "FFFFFF", True),
    "Skipped": ("D9D9D9", "3F3F3F", False),
    "Manual Application Submission Needed": ("FFEB9C", "9C6500", False),
}
RESPONSES = ("Interview", "Offer", "Rejected")


def _workbook(path=TRACKER_PATH):
    import openpyxl
    return openpyxl.load_workbook(path)


def initialize(path=TRACKER_PATH) -> dict:
    """Create the empty, policy-compliant tracker without overwriting an existing workbook."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    path = Path(path)
    if path.exists():
        try:
            openpyxl.load_workbook(path, read_only=True).close()
        except Exception as error:
            return {"ok": False, "created": False,
                    "error": f"tracker exists but is unreadable; left untouched: {error}"}
        return {"ok": True, "created": False, "path": str(path)}

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Applications"
    sheet.append(COLUMNS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L1"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = (14, 26, 38, 22, 48, 28, 14, 12, 22, 38, 60, 12)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    for row in range(2, 1001):
        sheet.cell(row=row, column=COLUMNS.index("Notes") + 1).alignment = Alignment(
            wrap_text=True, vertical="top"
        )

    validation_values = ",".join(STATUS_COLORS)
    validation = DataValidation(
        type="list", formula1=f'"{validation_values}"', allow_blank=True
    )
    sheet.add_data_validation(validation)
    validation.add("J2:J1000")

    legend = workbook.create_sheet("Legend")
    legend.append(("Status", "Fill", "Font", "Bold"))
    for cell in legend[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    for status, (fill, font, bold) in STATUS_COLORS.items():
        legend.append((status, fill, font, bold))
        cell = legend.cell(row=legend.max_row, column=1)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(color=font, bold=bold)
    for column, width in (("A", 42), ("B", 12), ("C", 12), ("D", 10)):
        legend.column_dimensions[column].width = width

    temporary = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb", suffix=".xlsx", prefix=f".{path.stem}.",
            dir=path.parent, delete=False,
        ) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        if path.exists():
            return {"ok": True, "created": False, "path": str(path)}
        os.replace(temporary, path)
        return {"ok": True, "created": True, "path": str(path)}
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()


def _rows(workbook) -> list[dict]:
    sheet = workbook["Applications"]
    out = []
    for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[1]:
            continue
        out.append({"row": index, **{name: row[i] for i, name in enumerate(COLUMNS) if i < len(row)}})
    return out


def _as_date(value):
    if hasattr(value, "date"):
        return value.date()
    if isinstance(value, (int, float)):
        return date(1899, 12, 30) + timedelta(days=int(value))
    if isinstance(value, str):
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            return None
    return None


def _matches(row: dict, company: str, job_id=None, url=None, title=None) -> bool:
    if norm(company) not in norm(row.get("Company")):
        return False
    if job_id and norm(job_id) != norm(row.get("Job ID")):
        return False
    if url and norm(url) != norm(row.get("Job URL")):
        return False
    if title and norm(title) != norm(row.get("Job Title")):
        return False
    return True


def _style(sheet, row_index: int, status: str) -> None:
    from openpyxl.styles import Font, PatternFill
    fill, font, bold = STATUS_COLORS.get(status, ("FFFFFF", "000000", False))
    cell = sheet.cell(row=row_index, column=COLUMNS.index("Status") + 1)
    cell.fill = PatternFill("solid", start_color=fill, end_color=fill)
    cell.font = Font(color=font, bold=bold)


def has_applied(company: str, job_id=None, url=None, title=None, path=TRACKER_PATH) -> dict:
    for row in _rows(_workbook(path)):
        if _matches(row, company, job_id, url, title):
            return {"applied": True, "row": row["row"], "title": row.get("Job Title"),
                    "status": row.get("Status"), "date": str(row.get("Date Applied") or "")}
    return {"applied": False}


def add(company: str, title: str, path=TRACKER_PATH, **fields) -> dict:
    workbook = _workbook(path)
    sheet = workbook["Applications"]
    row_index = sheet.max_row + 1
    values = {
        "Date Applied": fields.get("date") or today(),
        "Company": company,
        "Job Title": title,
        "Job ID": fields.get("job_id"),
        "Job URL": fields.get("url"),
        "Location": fields.get("location"),
        "Posted Date": fields.get("posted"),
        "Fit Score": fields.get("fit"),
        "Portal/ATS": fields.get("portal"),
        "Status": fields.get("status") or "Applied",
        "Notes": fields.get("notes"),
        "Match %": fields.get("match"),
    }
    for index, name in enumerate(COLUMNS, start=1):
        sheet.cell(row=row_index, column=index, value=values[name])
    _style(sheet, row_index, values["Status"])
    workbook.save(path)
    return {"ok": True, "row": row_index, "company": company, "title": title,
            "status": values["Status"]}


def update(company: str, job_id=None, url=None, title=None, status=None, match=None,
           note=None, path=TRACKER_PATH) -> dict:
    workbook = _workbook(path)
    sheet = workbook["Applications"]
    changed = []
    for row in _rows(workbook):
        if not _matches(row, company, job_id, url, title):
            continue
        if status:
            sheet.cell(row=row["row"], column=COLUMNS.index("Status") + 1, value=status)
            _style(sheet, row["row"], status)
            changed.append("Status")
        if match is not None:
            sheet.cell(row=row["row"], column=COLUMNS.index("Match %") + 1, value=match)
            changed.append("Match %")
        if note:
            cell = sheet.cell(row=row["row"], column=COLUMNS.index("Notes") + 1)
            existing = (cell.value or "").strip()
            cell.value = f"{existing}\n{today()} {note}".strip()
            changed.append("Notes")
        workbook.save(path)
        return {"ok": True, "row": row["row"], "changed": changed}
    return {"ok": False, "error": "no matching row"}


def summary(path=TRACKER_PATH) -> dict:
    rows = _rows(_workbook(path))
    counts: dict[str, int] = {}
    responses = []
    followups = []
    for row in rows:
        status = str(row.get("Status") or "").strip()
        counts[status or "blank"] = counts.get(status or "blank", 0) + 1
        applied = _as_date(row.get("Date Applied"))
        if status in RESPONSES:
            responses.append({"company": row.get("Company"), "status": status,
                              "date": str(applied or "")})
        if status == "Applied" and applied and applied + timedelta(days=7) <= date.today():
            followups.append({"company": row.get("Company"), "applied": str(applied)})
    return {"total": len(rows), "by_status": counts,
            "responses": sorted(responses, key=lambda r: r["date"], reverse=True)[:5],
            "followups_due": followups[:10]}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = parser.add_subparsers(dest="command", required=True)

    initialize_command = sub.add_parser("init", help="create the empty tracker if missing")
    initialize_command.add_argument("--file", type=Path, default=TRACKER_PATH,
                                    help="tracker workbook path")

    check = sub.add_parser("has-applied", help="duplicate check")
    check.add_argument("--company", required=True)
    check.add_argument("--job-id")
    check.add_argument("--url")
    check.add_argument("--title")

    sub.add_parser("summary", help="status counts, recent responses, follow-ups due")

    create = sub.add_parser("add", help="append one application row")
    create.add_argument("--company", required=True)
    create.add_argument("--title", required=True)
    for flag in ("--job-id", "--url", "--location", "--posted", "--portal", "--fit",
                 "--status", "--notes", "--match", "--date"):
        create.add_argument(flag)

    edit = sub.add_parser("update", help="change status, match score, or notes")
    edit.add_argument("--company", required=True)
    for flag in ("--job-id", "--url", "--title", "--status", "--match", "--note"):
        edit.add_argument(flag)

    args = parser.parse_args(argv)
    if args.command == "init":
        result = initialize(args.file)
        emit(result)
        return 0 if result.get("ok") else 1
    if args.command == "has-applied":
        emit(has_applied(args.company, args.job_id, args.url, args.title))
        return 0
    if args.command == "summary":
        emit(summary())
        return 0
    if args.command == "add":
        emit(add(
            args.company, args.title, date=args.date, job_id=args.job_id, url=args.url,
            location=args.location, posted=args.posted, fit=args.fit, portal=args.portal,
            status=args.status, notes=args.notes, match=args.match,
        ))
        return 0
    if not any((args.status, args.match, args.note)):
        print("nothing to update: pass --status, --match, or --note", file=sys.stderr)
        return 2
    result = update(args.company, args.job_id, args.url, args.title, args.status,
                    args.match, args.note)
    emit(result)
    return 0 if result.get("ok") else 1


if __name__ == "__main__":
    raise SystemExit(main())
