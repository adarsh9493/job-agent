"""Data-slice contract tests (Phase 2).

Every large data file must be reachable through a small JSON slice so a session never has to
read it whole. These tests pin the slice budgets and the read/write behaviour of each script.
"""
import contextlib
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
for sub in ("scan", "runtime", "tailoring", "maintenance"):
    sys.path.insert(0, str(ROOT / "scripts" / sub))

import answers
import companies
import mail
import queue
import state
import sweep
import tracker

AGENTS_BUDGET = 5120
SKILL_BUDGET = 5120
REFERENCE_BUDGET = 8192
SKILL_INDEX_BUDGET = 2048
HOT_BUDGET = 1024
QUEUE_BUDGET = 2048

BIG_STATE = {
    "mode": "running",
    "last_updated": "2026-10-31T09:00:00+00:00",
    "applications_this_session": 4,
    "session_stats": {"applied": 4, "submitted": 4, "skipped": 12, "failed": 0,
                      "manual_needed": 0, "queued": 1},
    "current": {"company": "Lucid Motors", "title": "ADAS Validation Engineer",
                "job_id": "4009378007", "step": "review", "folder": "Lucid_ADAS_2026-10-31"},
    "queue": [
        {"company": f"Company {index}", "title": "ADAS Validation Engineer " + "x" * 40,
         "job_id": str(index), "url": "https://example.com/role/" + str(index)}
        for index in range(1, 40)
    ],
    "sweep": {"cursor": "Hyundai America"},
    "interview_blocklist": {"active_interviews": [
        {"company": "Apptronik"}, {"company": "Wayve"},
    ]},
    "needs_user": ["a", "b", "c"],
}


def captured(func, *args, **kwargs):
    out = io.StringIO()
    with contextlib.redirect_stdout(out):
        code = func(*args, **kwargs)
    return code, out.getvalue()


class StateSliceTests(unittest.TestCase):
    def test_hot_projection_fits_the_budget(self):
        payload = state.fit(state.hot_payload(BIG_STATE))
        self.assertLessEqual(len(state.emit_text(payload).encode("utf-8")), HOT_BUDGET)
        self.assertEqual(payload["mode"], "running")
        self.assertEqual(payload["cursor"], "Hyundai America")
        self.assertLessEqual(payload["queue_len"], 40)

    def test_fit_trims_the_least_important_fields_first(self):
        payload = state.fit({"mode": "running", "updated": "u" * 900, "cursor": "c" * 400,
                             "queue": [{"company": "A", "title": "B"}], "queue_len": 1},
                            limit=600)
        self.assertLessEqual(len(state.emit_text(payload).encode("utf-8")), 600)
        self.assertNotIn("updated", payload)

    def test_dot_paths_read_and_write_nested_values(self):
        payload = {"sweep": {"cursor": "Waymo"}, "queue": [{"company": "Lucid"}]}
        self.assertEqual(state.dot_get(payload, "sweep.cursor"), "Waymo")
        self.assertEqual(state.dot_get(payload, "queue.0.company"), "Lucid")
        state.dot_set(payload, "sweep.cursor", "Torc")
        self.assertEqual(payload["sweep"]["cursor"], "Torc")
        with self.assertRaises(KeyError):
            state.dot_get(payload, "missing.path")

    def test_cli_writes_hot_json_within_budget(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            state_path = Path(temp_dir) / "state.json"
            hot_path = Path(temp_dir) / "hot.json"
            state_path.write_text(json.dumps(BIG_STATE), encoding="utf-8")
            with patch.object(state, "STATE_PATH", state_path), \
                    patch.object(state, "HOT_PATH", hot_path):
                code, out = captured(state.main, ["hot"])
                self.assertEqual(code, 0)
                self.assertLessEqual(hot_path.stat().st_size, HOT_BUDGET)
                self.assertLessEqual(len(out.encode("utf-8")), HOT_BUDGET)

                code, out = captured(state.main, ["get", "sweep.cursor"])
                self.assertEqual(code, 0)
                self.assertEqual(json.loads(out), "Hyundai America")

                code, _ = captured(state.main, ["set", "current.step", "form"])
                self.assertEqual(code, 0)
                saved = json.loads(state_path.read_text(encoding="utf-8"))
                self.assertEqual(saved["current"]["step"], "form")


class AnswersSliceTests(unittest.TestCase):
    def test_add_then_find_returns_the_recorded_answer(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "answers-bank.json"
            with patch.object(answers, "ANSWERS_PATH", path):
                code, _ = captured(
                    answers.main,
                    ["add", "Do you now or will you in the future require sponsorship?", "Yes"],
                )
                self.assertEqual(code, 0)

                code, out = captured(
                    answers.main,
                    ["find", "Will you require visa sponsorship?"],
                )
                self.assertEqual(code, 0)
                matches = json.loads(out)["matches"]
                self.assertEqual(matches[0]["answer"], "Yes")
                self.assertGreater(matches[0]["score"], 0)

                code, out = captured(answers.main, ["count"])
                self.assertEqual(json.loads(out), {"answers": 1})

    def test_find_is_empty_when_the_bank_is_missing(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            with patch.object(answers, "ANSWERS_PATH", Path(temp_dir) / "missing.json"):
                self.assertEqual(answers.find("anything at all"), [])


class TrackerSliceTests(unittest.TestCase):
    def workbook(self, path):
        import openpyxl
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Applications"
        sheet.append(tracker.COLUMNS)
        book.save(path)
        return path

    def test_add_query_update_and_summary(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = self.workbook(Path(temp_dir) / "applications.xlsx")

            result = tracker.add("Lucid Motors", "ADAS Validation Engineer", path=path,
                                 job_id="4009378007", status="Applied", match="74")
            self.assertTrue(result["ok"])

            found = tracker.has_applied("Lucid Motors", job_id="4009378007", path=path)
            self.assertTrue(found["applied"])
            self.assertEqual(found["status"], "Applied")
            self.assertFalse(tracker.has_applied("Lucid Motors", job_id="999", path=path)["applied"])

            updated = tracker.update("Lucid Motors", job_id="4009378007", status="Interview",
                                     match=81, note="Recruiter screen booked", path=path)
            self.assertTrue(updated["ok"])
            self.assertEqual(set(updated["changed"]), {"Status", "Match %", "Notes"})

            summary = tracker.summary(path=path)
            self.assertEqual(summary["total"], 1)
            self.assertEqual(summary["by_status"], {"Interview": 1})
            self.assertEqual(summary["responses"][0]["status"], "Interview")

    def test_init_creates_the_required_tracker_and_is_non_overwriting(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "applications.xlsx"
            result = tracker.initialize(path)
            self.assertTrue(result["ok"])
            self.assertTrue(result["created"])

            import openpyxl
            workbook = openpyxl.load_workbook(path)
            self.assertEqual(workbook.sheetnames, ["Applications", "Legend"])
            sheet = workbook["Applications"]
            self.assertEqual(tuple(cell.value for cell in sheet[1]), tracker.COLUMNS)
            self.assertTrue(any("J2:J1000" in str(rule.sqref)
                                for rule in sheet.data_validations.dataValidation))
            self.assertEqual(tuple(cell.value for cell in workbook["Legend"][1]),
                             ("Status", "Fill", "Font", "Bold"))
            workbook.close()

            original = path.read_bytes()
            again = tracker.initialize(path)
            self.assertTrue(again["ok"])
            self.assertFalse(again["created"])
            self.assertEqual(path.read_bytes(), original)

    def test_update_without_a_match_reports_cleanly(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = self.workbook(Path(temp_dir) / "applications.xlsx")
            result = tracker.update("Nobody", status="Applied", path=path)
            self.assertFalse(result["ok"])
            self.assertIn("no matching row", result["error"])


COMPANIES_FIXTURE = """# Target Companies (test fixture)

| # | Company | Sector | Careers URL | ATS |
|---|---|---|---|---|
| 1 | Target Automotive | Automotive | https://example.com/careers | Workday |
| 2 | Other Motors | Automotive | https://example.org/jobs | Greenhouse |

## Industry classification

| Company | Evidence | Business line | Verdict | Checked |
|---|---|---|---|---|
| Target Automotive | ADAS programs | OEM | IN | 2026-10-01 |
| Other Motors | EV only | OEM | IN | 2025-01-01 |
"""


class CompaniesSliceTests(unittest.TestCase):
    def test_build_show_and_stale(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "companies.md"
            source.write_text(COMPANIES_FIXTURE, encoding="utf-8")
            index = Path(temp_dir) / "companies.json"

            with patch.object(companies, "COMPANIES_JSON", index):
                result = companies.build(path=str(source), out=str(index))
                self.assertEqual(result["companies"], 2)
                self.assertEqual(result["classified"], 2)

                fresh = companies.show("Target Automotive")
                self.assertFalse(fresh["stale"])
                self.assertEqual(fresh["classification"]["verdict"], "IN")

                old = companies.show("Other Motors")
                self.assertTrue(old["stale"])

                self.assertEqual([row["name"] for row in companies.stale()], ["Other Motors"])
                self.assertIsNone(companies.show("Not Listed"))

    def test_missing_index_is_rebuilt_from_the_markdown(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "companies.md"
            source.write_text(COMPANIES_FIXTURE, encoding="utf-8")
            index = Path(temp_dir) / "companies.json"
            with patch.object(companies, "COMPANIES_JSON", index), \
                    patch.object(companies, "COMPANIES_MD", source):
                loaded = companies.load()
            self.assertEqual(len(loaded["companies"]), 2)
            self.assertTrue(index.is_file())


class MailSliceTests(unittest.TestCase):
    def test_query_is_narrow_and_policy_scoped(self):
        result = mail.build_query("LucidMotors.com", since="2026-10-29")
        self.assertIn("from:(@lucidmotors.com)", result["query"])
        self.assertIn("after:2026/10/29", result["query"])
        self.assertIn("only", result["scope"])

    def test_query_rejects_broad_or_invalid_input(self):
        with self.assertRaises(ValueError):
            mail.build_query("")
        with self.assertRaises(ValueError):
            mail.build_query("two domains.com")
        with self.assertRaises(ValueError):
            mail.build_query("lucidmotors.com", since="yesterday")


class QueueSliceTests(unittest.TestCase):
    def test_render_caps_markdown_at_the_budget(self):
        items = [
            {"company": f"Company {index}", "title": "ADAS Validation Engineer " + "y" * 60,
             "job_id": str(index)}
            for index in range(1, 30)
        ]
        text = queue.render(items, limit=300)
        self.assertLessEqual(len(text.encode("utf-8")), 300)
        self.assertIn("more", text)

    def test_render_lists_each_queued_role(self):
        text = queue.render([{"company": "Lucid", "title": "ADAS Validation Engineer",
                              "job_id": "1", "status": "queued"}])
        self.assertIn("Lucid", text)
        self.assertIn("ADAS Validation Engineer", text)
        self.assertLessEqual(len(text.encode("utf-8")), QUEUE_BUDGET)


class SweepNextTests(unittest.TestCase):
    def test_next_skips_interview_blocked_companies(self):
        entries = [
            {"company": "Apptronik", "last": None, "days": None, "idx": 1,
             "interview_blocked": True},
            {"company": "Lucid Motors", "last": None, "days": None, "idx": 2,
             "interview_blocked": False},
        ]
        with patch.object(sweep, "plan", return_value=entries), \
                patch.object(sys, "argv", ["sweep.py", "--next"]):
            code, out = captured(sweep.main)
        self.assertEqual(code, 0)
        self.assertEqual(json.loads(out)["next"], "Lucid Motors")

    def test_next_reports_when_every_company_is_blocked(self):
        entries = [{"company": "Wayve", "last": None, "days": None, "idx": 1,
                    "interview_blocked": True}]
        with patch.object(sweep, "plan", return_value=entries), \
                patch.object(sys, "argv", ["sweep.py", "--next"]):
            code, out = captured(sweep.main)
        self.assertEqual(code, 0)
        self.assertIsNone(json.loads(out)["next"])


class BudgetTests(unittest.TestCase):
    def test_agents_md_fits_the_session_budget(self):
        self.assertLessEqual((ROOT / "AGENTS.md").stat().st_size, AGENTS_BUDGET)

    def test_every_skill_and_reference_fits_its_budget(self):
        for path in sorted((ROOT / "agent" / "skills").rglob("*.md")):
            budget = SKILL_BUDGET if path.name == "SKILL.md" else REFERENCE_BUDGET
            with self.subTest(path=str(path.relative_to(ROOT))):
                self.assertLessEqual(path.stat().st_size, budget)

    def test_skill_frontmatter_index_fits_the_startup_budget(self):
        total = 0
        for folder in sorted(p for p in (ROOT / "agent" / "skills").iterdir() if p.is_dir()):
            text = (folder / "SKILL.md").read_text(encoding="utf-8")
            if text.startswith("---"):
                _, _, rest = text.partition("---")
                frontmatter, _, _ = rest.partition("---")
                total += len(frontmatter.encode("utf-8"))
        self.assertLessEqual(total, SKILL_INDEX_BUDGET)


if __name__ == "__main__":
    unittest.main()
